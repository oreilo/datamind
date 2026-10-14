from io import BytesIO
import json
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Preformatted
from reportlab.lib.styles import getSampleStyleSheet
from docx import Document
from html import escape

def render_report(title, payload, fmt):
    body=json.dumps(payload,indent=2,ensure_ascii=False,default=str)
    if fmt=="json":return body.encode("utf-8"),"application/json"
    if fmt=="html":return f"<html><head><meta charset='utf-8'><title>{escape(title)}</title></head><body><h1>{escape(title)}</h1><pre>{escape(body)}</pre></body></html>".encode(),"text/html"
    if fmt=="md":return (f"# {title}\n\n```json\n{body}\n```\n").encode(),"text/markdown"
    if fmt=="docx":
        doc=Document();doc.add_heading(title,0)
        for line in body.splitlines():doc.add_paragraph(line)
        buf=BytesIO();doc.save(buf);return buf.getvalue(),"application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    if fmt=="pdf":
        buf=BytesIO();styles=getSampleStyleSheet();story=[Paragraph(escape(title),styles['Title']),Spacer(1,12)]
        for i in range(0,len(body),3500):story.append(Preformatted(body[i:i+3500],styles['Code']))
        SimpleDocTemplate(buf).build(story);return buf.getvalue(),"application/pdf"
    raise ValueError("Unsupported format")
